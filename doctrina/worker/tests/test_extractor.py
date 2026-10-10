"""Pruebas unitarias de extracción; ejecutar con python -m unittest discover -s doctrina/worker/tests."""
import hashlib
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from extractor import ExtractionError, extract

class ExtractorTests(unittest.TestCase):
    def test_txt_preserva_texto_y_huella(self):
        raw = "Artículo 1. Deber\nContenido exacto.\n\nArtículo 2. Excepción\nSalvedad.".encode()
        result = extract(raw, "norma.txt")
        self.assertEqual(result["sha256"], hashlib.sha256(raw).hexdigest())
        self.assertEqual(result["cantidad_unidades"], 2)
        self.assertEqual(result["unidades"][0]["texto_fiel"], "Artículo 1. Deber\nContenido exacto.")
        self.assertEqual(result["estado"], "pendiente_revision")

    def test_docx_respeta_orden_de_parrafos_y_tablas(self):
        from docx import Document
        from io import BytesIO
        documento = Document()
        documento.add_paragraph("Primero")
        tabla = documento.add_table(rows=1, cols=1)
        tabla.cell(0, 0).text = "Segundo"
        documento.add_paragraph("Tercero")
        archivo = BytesIO()
        documento.save(archivo)
        unidades = extract(archivo.getvalue(), "intercalado.docx")["unidades"]
        self.assertEqual([u["texto_fiel"] for u in unidades],
                         ["Primero", "Segundo", "Tercero"])
        self.assertEqual([u["orden"] for u in unidades], [1, 2, 3])

    def test_rechaza_zip_falso_con_firma_pk(self):
        with self.assertRaises(ExtractionError):
            extract(b"PK" + b"basura", "falso.docx")

    def test_rechaza_zip_sin_estructura_docx(self):
        from io import BytesIO
        from zipfile import ZipFile
        archivo = BytesIO()
        with ZipFile(archivo, "w") as z:
            z.writestr("archivo.txt", "texto")
        with self.assertRaises(ExtractionError):
            extract(archivo.getvalue(), "incompleto.docx")

    def test_rechaza_zip_bomb_sintetico(self):
        from io import BytesIO
        from zipfile import ZipFile, ZIP_DEFLATED
        archivo = BytesIO()
        with ZipFile(archivo, "w", compression=ZIP_DEFLATED) as z:
            z.writestr("[Content_Types].xml", "contenido")
            z.writestr("word/document.xml", "A" * (2 * 1024 * 1024))
        with self.assertRaises(ExtractionError):
            extract(archivo.getvalue(), "compresion.docx")

    def test_rechaza_pdf_falso(self):
        with self.assertRaises(ExtractionError):
            extract(b"no es PDF", "falso.pdf")

    def test_rechaza_extension_peligrosa(self):
        with self.assertRaises(ExtractionError):
            extract(b"contenido", "ejecutar.exe")

    def test_rechaza_txt_binario(self):
        with self.assertRaises(ExtractionError):
            extract(b"abc\x00def", "binario.txt")

    def test_rechaza_archivo_vacio(self):
        with self.assertRaises(ExtractionError):
            extract(b"", "vacio.txt")

if __name__ == "__main__":
    unittest.main()
