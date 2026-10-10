"""Extracción conservadora para el núcleo privado Doctrina.

No invoca LLM, red, base de datos ni Storage. Solo recibe bytes desde el
backend autorizado y devuelve unidades con trazabilidad verificable.
"""
from __future__ import annotations

from dataclasses import dataclass, asdict
from hashlib import sha256
from io import BytesIO
from pathlib import PurePath
import re
from typing import Any

MAX_BYTES = 25 * 1024 * 1024
MAX_PAGES = 500
MAX_UNITS = 10000
MIMES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
}

class ExtractionError(ValueError):
    pass

@dataclass(frozen=True)
class Unit:
    orden: int
    tipo: str
    rotulo: str
    texto_fiel: str
    pagina_inicio: int | None
    pagina_fin: int | None
    huella: str

def _check(data: bytes, filename: str) -> str:
    if not data or len(data) > MAX_BYTES:
        raise ExtractionError("Archivo vacío o superior al límite")
    suffix = PurePath(filename).suffix.lower()
    if suffix not in MIMES:
        raise ExtractionError("Extensión no admitida")
    if suffix == ".pdf" and not data.startswith(b"%PDF-"):
        raise ExtractionError("Firma PDF inválida")
    if suffix == ".docx" and not data.startswith(b"PK"):
        raise ExtractionError("Firma DOCX inválida")
    if suffix == ".txt" and b"\x00" in data[:4096]:
        raise ExtractionError("TXT contiene bytes nulos")
    return suffix

def _unit(text: str, page: int | None, order: int) -> Unit:
    # No se normaliza el contenido: la fuente fiel prevalece sobre la estética.
    label_match = re.match(r"^\s*((?:Art(?:ículo|iculo)?\.?|Cap(?:ítulo|itulo)?\.?)\s*[\dIVXLCDM]+)", text, re.I)
    label = label_match.group(1) if label_match else ""
    return Unit(order, "articulo" if label.lower().startswith("art") else "bloque",
                label, text, page, page, sha256(text.encode("utf-8")).hexdigest())

def extract(data: bytes, filename: str) -> dict[str, Any]:
    suffix = _check(data, filename)
    units: list[Unit] = []
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(data), strict=True)
            if reader.is_encrypted:
                raise ExtractionError("PDF cifrado: revisión manual requerida")
            if len(reader.pages) > MAX_PAGES:
                raise ExtractionError("Excede máximo de páginas")
            for i, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                if not text.strip():
                    raise ExtractionError(f"Página {i} sin texto extraíble: requiere OCR supervisado")
                units.append(_unit(text, i, len(units) + 1))
        except ExtractionError:
            raise
        except Exception as exc:
            raise ExtractionError("No se pudo extraer PDF; requiere revisión") from exc
    elif suffix == ".docx":
        try:
            from docx import Document
            doc = Document(BytesIO(data))
            for para in doc.paragraphs:
                if para.text.strip():
                    units.append(_unit(para.text, None, len(units) + 1))
            for table in doc.tables:
                for row in table.rows:
                    text = "\t".join(cell.text for cell in row.cells)
                    if text.strip():
                        units.append(_unit(text, None, len(units) + 1))
        except Exception as exc:
            raise ExtractionError("No se pudo extraer DOCX; requiere revisión") from exc
    else:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeError as exc:
            raise ExtractionError("TXT debe estar codificado en UTF-8") from exc
        units.extend(_unit(block, None, i + 1)
                     for i, block in enumerate(re.split(r"\n\s*\n", text))
                     if block.strip())
    if not units or len(units) > MAX_UNITS:
        raise ExtractionError("Sin unidades válidas o demasiadas unidades")
    return {
        "sha256": sha256(data).hexdigest(),
        "mime_esperado": MIMES[suffix],
        "cantidad_unidades": len(units),
        "unidades": [asdict(u) for u in units],
        "estado": "pendiente_revision",
    }
